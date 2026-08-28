from pathlib import Path
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

FILES = [
    Path(r"C:/Users/PC/Downloads/Mogu_BA_Phan_he_Mon_an_v1.0.docx"),
    Path(r"C:/Users/PC/Downloads/MOGU_BA_003_Trang_chu_Dieu_huong_v1.0.docx"),
    Path(r"C:/Users/PC/Downloads/MOGU_BA_002_Onboarding_v1.0.docx"),
    Path(r"C:/Users/PC/Downloads/MOGU_BA_001_Dang_nhap_Dang_ky_v1.0.docx"),
]

for path in FILES:
    doc = Document(path)
    print(f"\n=== {path.name} | paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} ===")
    for p in doc.paragraphs:
        text = p.text.strip()
        if text and (p.style.name.startswith("Heading") or len(text) < 120):
            print(f"[{p.style.name}] {text}")
    for i, table in enumerate(doc.tables):
        rows = [" | ".join(c.text.strip().replace("\n", " / ") for c in row.cells) for row in table.rows[:4]]
        print(f"[TABLE {i+1} {len(table.rows)}x{len(table.columns)}] " + " || ".join(rows))
