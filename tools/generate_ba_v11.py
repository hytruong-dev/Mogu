from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"C:\QuangHy\Mogu\docs\ba")
OUT.mkdir(parents=True, exist_ok=True)

SOURCES = {
    "001": Path(r"C:\Users\PC\Downloads\MOGU_BA_001_Dang_nhap_Dang_ky_v1.0.docx"),
    "002": Path(r"C:\Users\PC\Downloads\MOGU_BA_002_Onboarding_v1.0.docx"),
    "003": Path(r"C:\Users\PC\Downloads\MOGU_BA_003_Trang_chu_Dieu_huong_v1.0.docx"),
    "004": Path(r"C:\Users\PC\Downloads\Mogu_BA_Phan_he_Mon_an_v1.0.docx"),
}

YELLOW = "FFD54F"
DARK = "161616"
CREAM = "FFF9E8"
GREEN = "24A865"
ORANGE = "F4A524"
RED = "E5484D"
GRAY = "626262"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=DARK, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(tbl_header)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8)
        shade(table.rows[0].cells[i], YELLOW)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=7.7)
            if str(value).startswith("THIẾU") or str(value).startswith("LỆCH"):
                shade(cells[i], "FDE8E8")
            elif str(value).startswith("MỘT PHẦN"):
                shade(cells[i], "FFF3D6")
            elif str(value).startswith("ĐÃ CÓ"):
                shade(cells[i], "E6F6EA")
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.name = "Arial"
        r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(3)


def add_intro(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(DARK)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor.from_string(GRAY)


def configure(doc, landscape=False):
    sec = doc.sections[-1]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)


def append_common(doc, ba_id, title):
    doc.add_page_break()
    add_intro(doc, f"PHỤ LỤC CẬP NHẬT {ba_id} — v1.1", title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Ngày đối chiếu: 13/08/2026 · Phạm vi: mobile, admin, backend và tài liệu API hiện hành")
    add_heading(doc, "1. Quy ước trạng thái", 1)
    add_table(doc, ["Trạng thái", "Ý nghĩa"], [
        ("ĐÃ CÓ", "Đã có mã nguồn/API và có dấu hiệu được nối vào giao diện."),
        ("MỘT PHẦN", "Đã có một phần UI, API hoặc nghiệp vụ nhưng chưa khép kín luồng."),
        ("THIẾU", "Chưa tìm thấy triển khai tương ứng trong mã nguồn hiện tại."),
        ("LỆCH TÀI LIỆU", "Tài liệu, API và/hoặc giao diện đang mô tả khác nhau."),
    ], [3, 14])


def clone_and_append(source, output, ba_id, title, sections):
    doc = Document(source)
    configure(doc)
    append_common(doc, ba_id, title)
    for heading, content in sections:
        add_heading(doc, heading, 1)
        if isinstance(content, tuple) and content[0] == "table":
            _, headers, rows, widths = content
            add_table(doc, headers, rows, widths)
        else:
            add_bullets(doc, content)
    doc.save(output)


def build_ba001():
    rows = [
        ("AUTH-01", "Đăng ký email/mật khẩu", "ĐÃ CÓ", "Backend và mobile có register; cần chuẩn hóa họ tên và xác nhận mật khẩu.", "P0"),
        ("AUTH-02", "Đăng nhập, refresh, me, logout", "ĐÃ CÓ", "API và mobile service đã tồn tại; cần kiểm thử refresh đồng thời và hết hạn phiên.", "P0"),
        ("AUTH-03", "Xác minh OTP / gửi lại OTP", "LỆCH TÀI LIỆU", "API_AUTH.md mô tả nhưng controller backend chưa công bố endpoint tương ứng.", "P0"),
        ("AUTH-04", "Quên/đặt lại mật khẩu", "MỘT PHẦN", "Backend có endpoint; cần xác nhận UI, deep link và trạng thái token hết hạn.", "P0"),
        ("AUTH-05", "Google/Apple", "MỘT PHẦN", "Có nút giao diện nhưng chưa có bằng chứng luồng OAuth hoàn chỉnh end-to-end.", "P1"),
        ("AUTH-06", "Bảo mật phiên", "MỘT PHẦN", "Cần secure storage, revoke theo thiết bị, rate limit, audit và chống refresh race.", "P0"),
    ]
    sections = [
        ("2. Ma trận đối chiếu", ("table", ["ID", "Yêu cầu", "Trạng thái", "Bằng chứng / khoảng trống", "Ưu tiên"], rows, [2, 4, 3, 8, 2])),
        ("3. Yêu cầu bổ sung bắt buộc", [
            "BA-001-ADD-01: Chọn một nguồn sự thật cho xác minh email: hoặc bổ sung POST /auth/verify-otp và POST /auth/resend-otp, hoặc xóa luồng OTP khỏi BA/API.",
            "BA-001-ADD-02: Hợp đồng đăng ký phải thống nhất các trường fullName, email, password, passwordConfirmation và consentVersion giữa UI, DTO và DB.",
            "BA-001-ADD-03: Access token chỉ lưu trong bộ nhớ; refresh token lưu bằng cơ chế bảo mật nền tảng, có rotation và revoke khi logout.",
            "BA-001-ADD-04: Mọi lỗi auth dùng một error envelope chung gồm code, message, fieldErrors, traceId.",
            "BA-001-ADD-05: Bổ sung trạng thái loading, offline, tài khoản bị khóa, email chưa xác minh, token hết hạn và retry có giới hạn.",
        ]),
        ("4. Tiêu chí nghiệm thu bổ sung", [
            "Đăng ký thành công chỉ tạo một tài khoản khi người dùng nhấn nhiều lần hoặc mạng chập chờn.",
            "Phiên được khôi phục sau khi mở lại ứng dụng; refresh lỗi đưa về đăng nhập và xóa dữ liệu nhạy cảm.",
            "Toàn bộ endpoint được kiểm tra contract tự động giữa mobile và backend.",
            "Google/Apple chỉ được đánh dấu hoàn thành khi đăng nhập thật, liên kết tài khoản và hủy liên kết đều hoạt động.",
        ]),
    ]
    clone_and_append(SOURCES["001"], OUT / "MOGU_BA_001_Dang_nhap_Dang_ky_v1.1.docx", "BA-001", "Đăng nhập & Đăng ký — đối chiếu triển khai", sections)


def build_ba002():
    rows = [
        ("ONB-01", "Khởi tạo/tiếp tục phiên onboarding", "ĐÃ CÓ", "Backend có start, get session và step endpoints.", "P0"),
        ("ONB-02", "Số bước", "LỆCH TÀI LIỆU", "API_ONBOARDING mô tả 8 bước; mobile hiện theo thiết kế rút gọn 6 bước.", "P0"),
        ("ONB-03", "Ngày sinh", "ĐÃ CÓ", "Có UI chọn ngày và API step; cần thống nhất timezone và tuổi tối thiểu.", "P0"),
        ("ONB-04", "Chiều cao + cân nặng", "MỘT PHẦN", "UI đã gộp; backend/tài liệu cũ còn đánh số bước khác.", "P0"),
        ("ONB-05", "Mục tiêu", "ĐÃ CÓ", "Catalog goals và profile/onboarding đã tồn tại.", "P0"),
        ("ONB-06", "Khẩu vị/chế độ ăn/dị ứng", "MỘT PHẦN", "Có UI và catalog, cần xác minh ghi nhận N-N và loại trừ khi random.", "P0"),
        ("ONB-07", "Tóm tắt/hoàn tất", "ĐÃ CÓ", "Có summary/complete; cần idempotency và điều hướng duy nhất.", "P0"),
    ]
    sections = [
        ("2. Quyết định chuẩn hóa luồng", [
            "BA v1.1 xác lập 6 bước hiển thị: (1) Chào mừng, (2) Ngày sinh, (3) Cơ thể — chiều cao và cân nặng, (4) Mục tiêu, (5) Ghi nhớ — khẩu vị/chế độ ăn/dị ứng, (6) Tóm tắt và hoàn tất.",
            "Tên và giới tính không còn là màn độc lập: tên lấy từ đăng ký/hồ sơ; giới tính là trường tùy chọn trong hồ sơ hoặc bottom sheet bổ sung.",
            "Backend phải cung cấp mapping tương thích từ step legacy 1–8 sang schema v2 1–6; không để mobile tự suy đoán số bước.",
        ]),
        ("3. Ma trận đối chiếu", ("table", ["ID", "Yêu cầu", "Trạng thái", "Bằng chứng / khoảng trống", "Ưu tiên"], rows, [2, 4, 3, 8, 2])),
        ("4. Yêu cầu bổ sung bắt buộc", [
            "BA-002-ADD-01: GET onboarding/session trả flowVersion, currentStep, completedSteps và dữ liệu đã lưu để resume.",
            "BA-002-ADD-02: PUT step dùng idempotency key hoặc version để tránh ghi đè khi retry.",
            "BA-002-ADD-03: Dị ứng phải là dữ liệu an toàn bắt buộc, không chỉ sở thích; kết quả random phải loại trừ tuyệt đối.",
            "BA-002-ADD-04: Skip phải được định nghĩa riêng cho từng bước và lưu lý do/null rõ ràng.",
            "BA-002-ADD-05: Hoàn tất cập nhật profile, đánh dấu onboardingCompleted và điều hướng Home đúng một lần.",
        ]),
        ("5. Tiêu chí nghiệm thu bổ sung", [
            "Đóng ứng dụng tại bất kỳ bước nào và mở lại phải trở về đúng bước cùng dữ liệu trước đó.",
            "Mobile chỉ hiển thị 6 chấm tiến trình và không còn bước cân nặng/chiều cao riêng lẻ.",
            "Không có biểu tượng check trên thẻ được chọn nếu thiết kế yêu cầu chỉ dùng nền vàng.",
            "Ảnh linh vật hiển thị đúng tỷ lệ trên Android/iOS và không bị phóng vượt khung khi NativeWind tải lỗi.",
        ]),
    ]
    clone_and_append(SOURCES["002"], OUT / "MOGU_BA_002_Onboarding_v1.1.docx", "BA-002", "Onboarding — chuẩn hóa luồng 6 bước", sections)


def build_ba003():
    rows = [
        ("NAV-01", "Open app → auth/onboarding/home", "ĐÃ CÓ", "App.tsx kiểm tra phiên và điều hướng theo trạng thái.", "P0"),
        ("NAV-02", "Home API", "ĐÃ CÓ", "Backend /home và mobile Home service đã tồn tại.", "P0"),
        ("NAV-03", "Navigation chuẩn", "THIẾU", "Mobile đang dùng useState route thủ công, chưa có stack/deep link/back-state chuẩn.", "P0"),
        ("NAV-04", "5 tab chính", "MỘT PHẦN", "UI có Home/Explore/Random/Health/Profile nhưng nhiều luồng là local/static.", "P0"),
        ("NAV-05", "Notification", "ĐÃ CÓ", "Có list/unread/read APIs và mobile service.", "P1"),
        ("NAV-06", "Offline/cache/error", "MỘT PHẦN", "Chưa có chuẩn chung cho stale cache, retry, empty/error state.", "P1"),
        ("NAV-07", "Analytics/deep link", "THIẾU", "Chưa thấy triển khai đầy đủ theo BA.", "P1"),
    ]
    sections = [
        ("2. Ma trận đối chiếu", ("table", ["ID", "Yêu cầu", "Trạng thái", "Bằng chứng / khoảng trống", "Ưu tiên"], rows, [2, 4, 3, 8, 2])),
        ("3. Yêu cầu bổ sung bắt buộc", [
            "BA-003-ADD-01: Thay bộ định tuyến useState bằng React Navigation hoặc Expo Router, có AuthStack, OnboardingStack và MainTabs.",
            "BA-003-ADD-02: Định nghĩa route contract và deep link cho dish, article, community post, notification và health day.",
            "BA-003-ADD-03: Mỗi màn có loading skeleton, empty, offline, stale-data và error/retry nhất quán.",
            "BA-003-ADD-04: Navbar là component dùng chung, safe-area đúng, active state duy nhất và không che nội dung.",
            "BA-003-ADD-05: Home, Explore, Random, Health, Profile phải tách data hooks khỏi component trình bày để kiểm thử.",
        ]),
        ("4. Tiêu chí nghiệm thu bổ sung", [
            "Nút Back hệ điều hành Android quay lại đúng màn, không thoát ứng dụng sai luồng.",
            "Deep link mở đúng màn sau đăng nhập và quay lại nguồn hợp lệ.",
            "Tất cả tab giữ trạng thái cuộn theo phiên và làm mới có chủ đích.",
            "Không có nội dung bị navbar che trên màn hình nhỏ hoặc thiết bị có safe area khác nhau.",
        ]),
    ]
    clone_and_append(SOURCES["003"], OUT / "MOGU_BA_003_Trang_chu_Dieu_huong_v1.1.docx", "BA-003", "Trang chủ & Điều hướng — đối chiếu triển khai", sections)


def build_ba004():
    rows = [
        ("FOOD-01", "Danh mục/dish public/admin CRUD", "ĐÃ CÓ", "Backend controllers và admin Foods page/service đã tồn tại.", "P0"),
        ("FOOD-02", "Media presign/commit/approve", "ĐÃ CÓ", "Admin media API đã có; cần test quyền và orphan cleanup.", "P1"),
        ("FOOD-03", "Lifecycle/review/version", "ĐÃ CÓ", "Có review queue, lifecycle, versions, rollback; cần contract test.", "P0"),
        ("FOOD-04", "Import jobs", "MỘT PHẦN", "Có create/list/log/retry/cancel; connector và xử lý dữ liệu thực cần xác minh.", "P0"),
        ("FOOD-05", "Random/select/history", "ĐÃ CÓ", "Backend có randomizations, select và history; mobile tích hợp chưa khép kín.", "P0"),
        ("FOOD-06", "Explore/detail/save", "MỘT PHẦN", "Backend có public dish và saved dishes; nhiều UI mobile còn mock/static.", "P0"),
        ("FOOD-07", "Admin users", "THIẾU", "UsersPage ghi rõ chưa có backend endpoint quản trị người dùng.", "P1"),
        ("FOOD-08", "Admin community moderation", "THIẾU", "CommunityPage ghi rõ chưa có backend endpoint moderation.", "P1"),
        ("FOOD-09", "Admin reports", "THIẾU", "ReportsPage ghi rõ chưa có backend analytics endpoint.", "P1"),
        ("FOOD-10", "CORS và cấu hình môi trường", "MỘT PHẦN", "Ảnh lỗi hiện tại cho thấy request admin bị CORS; cần allowlist theo môi trường.", "P0"),
    ]
    sections = [
        ("2. Ma trận đối chiếu", ("table", ["ID", "Yêu cầu", "Trạng thái", "Bằng chứng / khoảng trống", "Ưu tiên"], rows, [2, 4, 3, 8, 2])),
        ("3. Yêu cầu bổ sung kiến trúc và dữ liệu", [
            "BA-004-ADD-01: Một DishesModule duy nhất; loại bỏ đường dẫn dishes-v2 và giải quyết dứt điểm enum/model category.",
            "BA-004-ADD-02: Chuẩn hóa Goal, Allergen, Ingredient, DietType và các bảng N-N; khai báo unique, index và delete policy.",
            "BA-004-ADD-03: Dish có version; endpoint cập nhật dùng If-Match để chống mất dữ liệu khi hai admin sửa đồng thời.",
            "BA-004-ADD-04: DishAuditLog ghi actor, action, before, after, reason, traceId và thời gian.",
            "BA-004-ADD-05: RBAC dựa trên ProfileRole; tách quyền viewer, editor, reviewer, publisher, admin.",
            "BA-004-ADD-06: Import job có progress, retry/cancel, logs, source evidence, conflict state và resume sau lỗi.",
            "BA-004-ADD-07: Random phải mô tả loại trừ dị ứng, scoring, fallback và lưu snapshot input/output để giải thích kết quả.",
            "BA-004-ADD-08: CORS dùng ADMIN_ORIGINS/MOBILE_ORIGINS theo môi trường; localhost:5175 phải được khai báo ở dev.",
        ]),
        ("4. API còn thiếu cho Admin", [
            "GET/PATCH /admin/users, GET /admin/users/:id, POST /admin/users/:id/restrict và export người dùng.",
            "GET /admin/community/reports, GET /admin/community/reports/:id, POST action ignore/hide/warn/delete.",
            "GET /admin/analytics/overview, random, retention, popular-needs, top-dishes và export report.",
            "Chuẩn hóa pagination/filter/sort/error envelope cho mọi danh sách admin.",
        ]),
        ("5. Tiêu chí nghiệm thu bổ sung", [
            "Admin localhost gọi API dev không phát sinh lỗi CORS và preflight OPTIONS trả đúng header.",
            "Một món đi trọn draft → review → published → archived và rollback, có audit đầy đủ.",
            "Import job lỗi có thể retry/cancel, log rõ nguồn và không tạo bản ghi trùng.",
            "Món chứa dị ứng đã khai báo không bao giờ xuất hiện trong kết quả random.",
            "Mobile Explore/Random/Detail/Save dùng dữ liệu backend, không phụ thuộc fixture trong bản production.",
        ]),
    ]
    clone_and_append(SOURCES["004"], OUT / "MOGU_BA_004_Phan_he_Mon_an_v1.1.docx", "BA-004", "Phân hệ Món ăn — đối chiếu Mobile, Admin và Backend", sections)


def build_matrix():
    doc = Document()
    configure(doc, landscape=True)
    add_intro(doc, "MOGU — MA TRẬN COVERAGE BA / MOBILE / ADMIN / BACKEND", "Phiên bản 1.0 · Ngày 13/08/2026")
    add_heading(doc, "1. Kết luận điều hành", 1)
    add_bullets(doc, [
        "Backend đã có nền tảng tốt cho auth, onboarding, home, taxonomy, dish lifecycle, review, import job, media và randomization.",
        "Khoảng trống lớn nhất là sự lệch contract: OTP được tài liệu hóa nhưng thiếu endpoint; onboarding 8 bước trong backend docs nhưng UI đã rút còn 6 bước; mobile dùng điều hướng thủ công.",
        "Admin có ba khu vực UI chưa có backend tương ứng: Người dùng, Cộng đồng và Báo cáo.",
        "Nhiều màn mobile đã hoàn thiện về UI nhưng chưa có bằng chứng tích hợp API end-to-end; cần tách rõ mock/demo và production data.",
        "Ưu tiên trước phát triển thêm UI: khóa OpenAPI, CORS/env, auth/onboarding contract, navigation và contract tests.",
    ])
    add_heading(doc, "2. Ma trận tổng hợp", 1)
    rows = [
        ("BA-001", "Auth", "MỘT PHẦN", "Không áp dụng", "ĐÃ CÓ", "OTP, social auth, secure session và contract register."),
        ("BA-002", "Onboarding", "MỘT PHẦN", "Không áp dụng", "ĐÃ CÓ", "Chuẩn hóa 6 bước và mapping legacy 8 bước."),
        ("BA-003", "Home & Navigation", "MỘT PHẦN", "Không áp dụng", "ĐÃ CÓ", "Router chuẩn, deep link, cache và state restoration."),
        ("BA-004", "Dish public/random", "MỘT PHẦN", "MỘT PHẦN", "ĐÃ CÓ", "Nối API mobile; hoàn thiện import connectors."),
        ("BA-004A", "Admin users", "Không áp dụng", "THIẾU", "THIẾU", "User management endpoints và RBAC."),
        ("BA-004B", "Admin community", "MỘT PHẦN", "THIẾU", "THIẾU", "Moderation/report APIs và actions."),
        ("BA-004C", "Admin analytics", "MỘT PHẦN", "THIẾU", "THIẾU", "Analytics aggregates và export."),
    ]
    add_table(doc, ["BA", "Phạm vi", "Mobile", "Admin", "Backend", "Khoảng trống trọng yếu"], rows, [2.2, 4, 3, 3, 3, 10])
    add_heading(doc, "3. Backlog ưu tiên", 1)
    backlog = [
        ("P0-01", "P0", "Backend", "Sửa CORS/env allowlist cho admin dev/prod; kiểm thử preflight."),
        ("P0-02", "P0", "All", "Xuất OpenAPI chuẩn và tạo contract tests cho mobile/admin."),
        ("P0-03", "P0", "Auth", "Quyết định và hoàn tất OTP; chuẩn hóa register DTO."),
        ("P0-04", "P0", "Onboarding", "Chốt flowVersion 6 bước, migration/mapping và resume."),
        ("P0-05", "P0", "Mobile", "Chuyển sang navigation stack/tab chuẩn và deep link."),
        ("P0-06", "P0", "Food", "Nối Explore/Random/Detail/Save với backend và loại bỏ mock production."),
        ("P0-07", "P0", "Food", "Bảo đảm rule loại trừ dị ứng và audit random."),
        ("P1-01", "P1", "Admin", "Xây user management APIs và quyền hạn."),
        ("P1-02", "P1", "Admin", "Xây community moderation APIs."),
        ("P1-03", "P1", "Admin", "Xây analytics/report APIs và export."),
        ("P1-04", "P1", "Quality", "Visual regression cho thiết bị mobile và responsive admin."),
        ("P1-05", "P1", "Ops", "Observability: traceId, metrics, audit, alert và retention."),
        ("P2-01", "P2", "Data", "pgvector/recommendation nâng cao sau khi dữ liệu chuẩn ổn định."),
    ]
    add_table(doc, ["ID", "Ưu tiên", "Phân hệ", "Hạng mục"], backlog, [2.5, 2, 3.5, 16])
    add_heading(doc, "4. Quy tắc hoàn tất chung (Definition of Done)", 1)
    add_bullets(doc, [
        "Yêu cầu có BA ID, API contract, migration/seed, unit/integration/E2E test và tiêu chí nghiệm thu.",
        "Không đánh dấu hoàn thành khi chỉ có UI mock hoặc endpoint chưa được client sử dụng.",
        "Mọi API có auth/RBAC, validation, pagination, error envelope, traceId và audit phù hợp.",
        "Mobile được kiểm tra Android/iOS, safe area, bàn phím, offline và kích thước màn hình nhỏ.",
        "Admin được kiểm tra Chrome/Edge, responsive tối thiểu, CORS dev/staging/prod và quyền theo vai trò.",
        "Tài liệu API và BA được cập nhật trong cùng pull request với thay đổi mã nguồn.",
    ])
    add_heading(doc, "5. Tệp tham chiếu đã đối chiếu", 1)
    add_bullets(doc, [
        "Bốn tài liệu BA v1.0 do người dùng cung cấp.",
        "Backend controllers và API_AUTH.md, API_ONBOARDING.md, BA-003_API.md, BA-004_API.md.",
        "Mobile App.tsx, API services và các màn auth/onboarding/home/explore/random/health/profile.",
        "Admin routes, services và các trang Foods, Import, Review, Users, Community, Reports.",
        "Nội dung cuộc hội thoại tham chiếu về kiến trúc BA-004 và lộ trình ingestion có kiểm soát.",
    ])
    doc.save(OUT / "MOGU_BA_000_Ma_tran_coverage_Mobile_Admin_Backend_v1.0.docx")


if __name__ == "__main__":
    build_ba001()
    build_ba002()
    build_ba003()
    build_ba004()
    build_matrix()
    print("Generated:")
    for p in sorted(OUT.glob("*.docx")):
        print(p)
